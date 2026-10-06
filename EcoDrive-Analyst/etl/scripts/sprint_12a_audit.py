"""Sprint 12A source inventory and non-destructive grain audit.

Run from the repository root:
    python etl/scripts/sprint_12a_audit.py

The script reads only from etl/data/{raw,reference,evidence}.  Its machine-
readable findings are written below etl/data/processed/sprint_12a_audit/.
"""
from __future__ import annotations

import csv
import itertools
import json
import math
import os
import re
import statistics
import tempfile
import zipfile
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

import openpyxl
import pandas as pd


ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT / "etl" / "data"
OUT = DATA / "processed" / "sprint_12a_audit"
MAX_EXACT_UNIQUES = 1_000_000
EXAMPLE_COUNT = 5


def clean_value(value: Any) -> Any:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def display(value: Any) -> str:
    value = clean_value(value)
    if value is None:
        return ""
    text = str(value).replace("\n", " ").strip()
    return text[:160]


def source_family(path: Path) -> str:
    p = str(path).lower()
    if "epa_testcar" in p:
        return "EPA Test Car"
    if "epa_certification" in p and "models" in p:
        return "EPA Certified Models"
    if "epa_certification" in p:
        return "EPA Certified Test Results"
    if "fueleconomy" in p:
        return "FuelEconomy.gov"
    if "wltp_eea" in p:
        return "EEA WLTP / CO2 monitoring"
    if "wltp_jrc" in p:
        return "JRC technical vehicle dataset"
    if "evcis" in p:
        return "EPA EV-CIS reference"
    if "reference\\wltp" in p or "reference/wltp" in p:
        return "UNECE R154 reference"
    if "evidence" in p:
        return "Certification evidence PDF"
    return "Other"


def file_notes(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return "PDF evidence/reference; inventoried only (not parsed or used for RAG)."
    if suffix == ".zip":
        return "Archive; workbook/XML members inventoried without extraction to raw."
    if path.name == "data.csv":
        return "Large EEA CSV profiled in streaming mode; exact cardinalities are capped when >1,000,000."
    return "Structured source; read-only audit."


def detect_header(ws: Any, limit: int = 8) -> int | None:
    """Find the first row containing a recognizable tabular header."""
    for idx, row in enumerate(ws.iter_rows(min_row=1, max_row=limit, values_only=True), 1):
        values = [display(x) for x in row]
        joined = " | ".join(values).lower()
        if ("model year" in joined or "model yr" in joined or "oem anon" in joined) and sum(bool(x) for x in values) >= 3:
            return idx
    return None


def dtype_of(values: list[Any]) -> str:
    nonnull = [x for x in values if x is not None]
    if not nonnull:
        return "empty"
    if all(isinstance(x, bool) for x in nonnull):
        return "boolean"
    if all(isinstance(x, int) and not isinstance(x, bool) for x in nonnull):
        return "integer"
    if all(isinstance(x, (int, float)) and not isinstance(x, bool) for x in nonnull):
        return "number"
    if all(isinstance(x, (date, datetime)) for x in nonnull):
        return "date"
    if all(isinstance(x, str) for x in nonnull):
        return "string"
    return "mixed"


def engineering_meaning(name: str, family: str) -> tuple[str, str, str]:
    """Only mark meanings direct where label and values make the concept explicit."""
    n = name.lower()
    unit = ""
    for match in re.findall(r"\[([^\]]+)\]|\(([^)]+)\)", name):
        unit = next((x for x in match if x), unit)
    direct = [
        (("target coefficient a" in n), "Target road-load coefficient A"),
        (("target coefficient b" in n), "Target road-load coefficient B"),
        (("target coefficient c" in n), "Target road-load coefficient C"),
        (("set coefficient a" in n), "Set road-load coefficient A"),
        (("set coefficient b" in n), "Set road-load coefficient B"),
        (("set coefficient c" in n), "Set road-load coefficient C"),
        (("equivalent test weight" in n or "vehicle mass (wltp)" in n), "Explicit test mass / ETW"),
        (("curb" in n and "mass" in n) or "curb weight" in n, "Explicit curb mass/weight"),
        (("tyre code" in n or "tire code" in n), "Tire/tyre description"),
        (("n gears" in n or "# of gears" in n or "number of gears" in n), "Transmission gear count"),
        (("transmission type" in n or "gear box type" in n or n == "trans"), "Transmission type"),
        (("axle ratio" in n), "Axle/final-drive ratio"),
        (("n/v ratio" in n), "N/V ratio"),
        (("engine capacity" in n or "displacement" in n or "eng displ" in n), "Engine displacement"),
        (("engine max power" in n or "rated horsepower" in n), "Engine power"),
        (("electric motor" in n), "Electric motor information"),
        (("battery" in n), "Battery information"),
        (("co2" in n), "CO2 result"),
        (("fuel consumption" in n), "Fuel-consumption result"),
        (("electric consumption" in n), "Electric-consumption result"),
        (("electric range" in n or n.startswith("range")), "Range result"),
        (("test procedure" in n), "Test procedure"),
        (("test number" in n), "Test identifier"),
        (("testgroup" in n or "test group" in n), "Test group"),
        (("fuel type" in n or "fuel category" in n or "fuel usage" in n), "Fuel type/category"),
        (("hybrid" in n or "is_plugin" in n or "is_electric" in n), "Electrification flag"),
        ((n in {"id", "vfn", "mp", "mh", "man", "mms", "tan", "va", "ve", "mk"}), "Identity field (meaning unresolved without EEA data dictionary)"),
        (("model year" in n or n == "year"), "Model year / reporting year"),
        (("make" in n or "manufacturer" in n or "mfr name" in n or "division" in n), "Vehicle make/manufacturer"),
        (("model" in n or "carline" in n), "Vehicle model/carline"),
        (("drive" in n), "Driveline / drive configuration"),
    ]
    for condition, meaning in direct:
        if condition:
            confidence = "DIRECTLY OBSERVED" if "unresolved" not in meaning else "UNRESOLVED"
            return meaning, unit, confidence
    return "TENTATIVE: no engineering meaning established in this audit", unit, "UNRESOLVED"


def profile_dataframe(df: pd.DataFrame, source: str, sheet: str) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for col in df.columns:
        series = df[col]
        values = [clean_value(v) for v in series.dropna().head(2000).tolist()]
        examples = []
        for value in values:
            rendered = display(value)
            if rendered and rendered not in examples:
                examples.append(rendered)
            if len(examples) == EXAMPLE_COUNT:
                break
        meaning, unit, confidence = engineering_meaning(str(col), source)
        rows.append({
            "source": source,
            "sheet": sheet,
            "original_field_name": str(col),
            "dtype": str(series.dtype),
            "non_null_count": int(series.notna().sum()),
            "non_null_pct": round(float(series.notna().mean() * 100), 3),
            "unique_count": int(series.nunique(dropna=True)),
            "examples": " | ".join(examples),
            "tentative_engineering_meaning": meaning,
            "unit_if_explicit": unit,
            "interpretation_confidence": confidence,
            "reference_used": "None; interpretation is based on explicit source field label and observed values only.",
        })
    return rows


def profile_csv_stream(path: Path, source: str) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """Profile EEA exactly for row/null counts; cardinality is exact until a documented cap."""
    with path.open("r", encoding="utf-8-sig", newline="", errors="replace") as handle:
        reader = csv.reader(handle)
        headers = next(reader)
        nonnull = [0] * len(headers)
        examples = [[] for _ in headers]
        numeric = [True] * len(headers)
        booleans = [True] * len(headers)
        uniques: list[set[str] | None] = [set() for _ in headers]
        capped = [False] * len(headers)
        row_count = 0
        for row_count, row in enumerate(reader, 1):
            if len(row) < len(headers):
                row += [""] * (len(headers) - len(row))
            for i, value in enumerate(row[: len(headers)]):
                value = value.strip()
                if not value:
                    continue
                nonnull[i] += 1
                if len(examples[i]) < EXAMPLE_COUNT and value not in examples[i]:
                    examples[i].append(value[:160])
                if numeric[i]:
                    try:
                        float(value.replace(",", "."))
                    except ValueError:
                        numeric[i] = False
                if booleans[i] and value.lower() not in {"true", "false", "0", "1", "yes", "no", "y", "n"}:
                    booleans[i] = False
                if uniques[i] is not None:
                    uniques[i].add(value)
                    if len(uniques[i]) > MAX_EXACT_UNIQUES:
                        uniques[i] = None
                        capped[i] = True
        profile = []
        for i, header in enumerate(headers):
            meaning, unit, confidence = engineering_meaning(header, source)
            dtype = "number" if numeric[i] else "boolean" if booleans[i] else "string"
            unique = f">={MAX_EXACT_UNIQUES + 1} (capped)" if capped[i] else len(uniques[i] or [])
            profile.append({
                "source": source,
                "sheet": "CSV",
                "original_field_name": header,
                "dtype": dtype,
                "non_null_count": nonnull[i],
                "non_null_pct": round((nonnull[i] / row_count * 100) if row_count else 0, 3),
                "unique_count": unique,
                "examples": " | ".join(examples[i]),
                "tentative_engineering_meaning": meaning,
                "unit_if_explicit": unit,
                "interpretation_confidence": confidence,
                "reference_used": "None; direct label/value observation. Cardinality capped above 1,000,000 to keep streaming audit bounded.",
            })
    return profile, {"rows": row_count, "columns": len(headers)}


def load_excel(path: Path, sheet: str, header_row: int) -> pd.DataFrame:
    return pd.read_excel(path, sheet_name=sheet, header=header_row - 1, engine="openpyxl")


def load_fueleconomy(path: Path) -> dict[str, pd.DataFrame]:
    frames: dict[str, pd.DataFrame] = {}
    with zipfile.ZipFile(path) as archive:
        member = next(x for x in archive.namelist() if x.lower().endswith(".xlsx"))
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as temp:
            temp.write(archive.read(member))
            tmp = Path(temp.name)
    try:
        book = openpyxl.load_workbook(tmp, read_only=True, data_only=True)
        for ws in book.worksheets:
            row = detect_header(ws) or 1
            frames[ws.title] = pd.read_excel(tmp, sheet_name=ws.title, header=row - 1, engine="openpyxl")
        book.close()
    finally:
        tmp.unlink(missing_ok=True)
    return frames


def compact_tuple(row: pd.Series, fields: list[str]) -> tuple[Any, ...]:
    return tuple(clean_value(row.get(field)) for field in fields)


def present(value: Any) -> bool:
    return value is not None and not (isinstance(value, float) and math.isnan(value))


def epa_audit(df: pd.DataFrame) -> tuple[dict[str, Any], list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    epa = df[df["Model Year"].eq(2026)].copy()
    epa.insert(0, "source_row", epa.index + 2)
    group_fields = ["Model Year", "Represented Test Veh Make", "Represented Test Veh Model"]
    target_fields = ["Target Coef A (lbf)", "Target Coef B (lbf/mph)", "Target Coef C (lbf/mph**2)"]
    set_fields = ["Set Coef A (lbf)", "Set Coef B (lbf/mph)", "Set Coef C (lbf/mph**2)"]
    roadload_fields = target_fields
    vehicle_group_size = epa.groupby(group_fields, dropna=False)["source_row"].transform("size")
    epa["make_model_year_group_rows"] = vehicle_group_size
    epa["target_abc"] = epa.apply(lambda r: compact_tuple(r, target_fields), axis=1)
    epa["set_abc"] = epa.apply(lambda r: compact_tuple(r, set_fields), axis=1)
    epa["roadload_config"] = epa["target_abc"]
    epa["test_identifier"] = epa.apply(lambda r: " | ".join(display(r[x]) for x in ["Test Number", "Actual Tested Testgroup"]), axis=1)
    epa["config_proxy"] = epa.apply(lambda r: " | ".join(display(r[x]) for x in ["Test Vehicle ID", "Test Veh Configuration #", "Test Number"]), axis=1)

    all_target = epa[target_fields].notna().all(axis=1)
    component_fields = ["Equivalent Test Weight (lbs.)", "Tested Transmission Type", "# of Gears", "Axle Ratio", "N/V Ratio"]
    all_components = epa[component_fields].notna().all(axis=1)
    any_components = epa[component_fields].notna().any(axis=1)
    # There is no tire field in this source, so full component closure cannot be directly observed.
    epa["roadload_availability"] = "NO_AUTHORITATIVE_ROADLOAD"
    epa.loc[all_target & ~any_components, "roadload_availability"] = "ROADLOAD_WITHOUT_COMPONENT_DATA"
    epa.loc[all_target & any_components, "roadload_availability"] = "ROADLOAD_WITH_PARTIAL_COMPONENT_DATA"
    epa.loc[all_target & all_components & False, "roadload_availability"] = "ROADLOAD_AVAILABLE"

    config_rows = epa[["source_row", *group_fields, "Test Vehicle ID", "Test Veh Configuration #", "Actual Tested Testgroup", "Test Number", "Test Procedure Cd", "Test Procedure Description", "Equivalent Test Weight (lbs.)", "Tested Transmission Type", "# of Gears", "Drive System Description", "Axle Ratio", "N/V Ratio", *target_fields, *set_fields, "roadload_availability", "make_model_year_group_rows"]].copy()
    config_rows["target_abc"] = epa["target_abc"].map(str)
    config_rows["set_abc"] = epa["set_abc"].map(str)

    grouped = epa.groupby(group_fields, dropna=False)
    grain = {
        "source_rows": int(len(epa)),
        "make_model_year_groups": int(grouped.ngroups),
        "groups_with_multiple_source_rows": int((grouped.size() > 1).sum()),
        "unique_test_identifiers_or_groups": int(epa["test_identifier"].nunique()),
        "distinct_roadload_configurations": int(epa["roadload_config"].nunique()),
        "distinct_target_abc_sets": int(epa["target_abc"].nunique()),
        "distinct_set_abc_sets": int(epa["set_abc"].nunique()),
        "distinct_etw_variants": int(epa["Equivalent Test Weight (lbs.)"].nunique()),
        "distinct_cycle_test_procedure_variants": int(epa[["Test Procedure Cd", "Test Procedure Description"]].drop_duplicates().shape[0]),
        "groups_with_multiple_roadload_configurations": int(grouped["roadload_config"].nunique().gt(1).sum()),
    }

    stable = ["Model Year", "Represented Test Veh Make", "Represented Test Veh Model", "Actual Tested Testgroup", "Test Veh Displacement (L)", "Tested Transmission Type", "# of Gears", "Drive System Description", "Test Procedure Cd"]
    pairs: list[dict[str, Any]] = []
    capped = False
    def add_pairs(keys: list[str], changed: list[str], candidate: str, limit: int = 5000) -> None:
        nonlocal capped
        for key, part in epa.groupby(keys, dropna=False):
            variants = part.drop_duplicates(changed)
            if len(variants) < 2:
                continue
            for left, right in itertools.combinations(variants.to_dict("records"), 2):
                if len(pairs) >= limit:
                    capped = True
                    return
                left_d, right_d = left, right
                difference = []
                for field in changed:
                    if clean_value(left_d[field]) != clean_value(right_d[field]):
                        difference.append(f"{field}: {display(left_d[field])} -> {display(right_d[field])}")
                if not difference:
                    continue
                target_delta = []
                for f in target_fields:
                    a, b = left_d[f], right_d[f]
                    if present(a) and present(b): target_delta.append(f"{f}: {float(b)-float(a):.8g}")
                mass_delta = ""
                a, b = left_d["Equivalent Test Weight (lbs.)"], right_d["Equivalent Test Weight (lbs.)"]
                if present(a) and present(b): mass_delta = str(float(b) - float(a))
                pairs.append({
                    "pair_id": f"EPA26-{len(pairs)+1:05d}",
                    "candidate_type": candidate,
                    "source_row_left": int(left_d["source_row"]),
                    "source_row_right": int(right_d["source_row"]),
                    "vehicle_configuration_identifiers": " | ".join(display(left_d[x]) for x in ["Model Year", "Represented Test Veh Make", "Represented Test Veh Model", "Actual Tested Testgroup", "Test Vehicle ID", "Test Veh Configuration #"]),
                    "shared_fields": "; ".join(f"{x}={display(left_d[x])}" for x in keys),
                    "differing_fields": "; ".join(difference),
                    "target_abc_delta": "; ".join(target_delta),
                    "mass_delta_lbs": mass_delta,
                    "candidate_interpretation": "DIRECTLY OBSERVED candidate; no component causal delta is claimed.",
                    "confidence": "DIRECTLY OBSERVED",
                    "warnings": "Candidate discovery only. Grain and physical isolation remain unresolved; no tire field exists in EPA Test Car source.",
                })
    # Strictly match all listed context fields, then expose only the changing factor(s).
    add_pairs(stable + target_fields + set_fields, ["Equivalent Test Weight (lbs.)"], "same context + roadload; different ETW")
    add_pairs(stable + ["Equivalent Test Weight (lbs.)"], target_fields + set_fields, "same context + ETW; different roadload set")
    drive_keys = [x for x in stable if x != "Drive System Description"] + ["Equivalent Test Weight (lbs.)", *target_fields, *set_fields]
    add_pairs(drive_keys, ["Drive System Description"], "same context; different drive configuration")
    grain["pair_output_capped"] = capped
    grain["pair_candidates_written"] = len(pairs)

    coverage_fields = {
        "authoritative Target ABC": target_fields,
        "Set ABC": set_fields,
        "mass / ETW": ["Equivalent Test Weight (lbs.)"],
        "tire specification": [],
        "tire pressure": [],
        "RRC / RR information": [],
        "transmission": ["Tested Transmission Type"],
        "gear count": ["# of Gears"],
        "axle/final-drive ratio": ["Axle Ratio"],
        "N/V": ["N/V Ratio"],
        "Cd / CdA": [],
        "TOTAL / NET information": [],
    }
    coverage: list[dict[str, Any]] = []
    proxy = epa.groupby("config_proxy", dropna=False)
    for concept, fields in coverage_fields.items():
        if fields:
            mask = epa[fields].notna().all(axis=1)
            proxy_mask = proxy[fields].apply(lambda p: p.notna().all(axis=1).any())
        else:
            mask = pd.Series(False, index=epa.index)
            proxy_mask = pd.Series(False, dtype=bool)
        coverage.append({
            "source": "EPA Test Car MY2026",
            "concept": concept,
            "source_row_count": int(len(epa)),
            "source_rows_available": int(mask.sum()),
            "source_row_coverage_pct": round(float(mask.mean() * 100), 3),
            "candidate_configuration_proxy": "Test Vehicle ID + Configuration # + Test Number (not semantically confirmed)",
            "candidate_configurations": int(proxy.ngroups),
            "candidate_configurations_available": int(proxy_mask.sum()),
            "candidate_configuration_coverage_pct": round(float(proxy_mask.mean() * 100) if len(proxy_mask) else 0, 3),
            "evidence_standard": "DIRECTLY OBSERVED" if fields else "UNRESOLVED / ABSENT",
            "notes": " | ".join(fields) if fields else "No explicit field located in EPA Test Car source.",
        })
    return grain, config_rows.to_dict("records"), pairs, coverage


def wltp_coverage(eea_profile: list[dict[str, Any]], jrc_profile: list[dict[str, Any]], eea_meta: dict[str, Any], jrc: pd.DataFrame) -> list[dict[str, Any]]:
    rows = []
    concepts = ["identity", "mass", "power", "fuel consumption", "electric consumption", "CO2", "range", "tire", "roadload", "phase"]
    for source, profile, total in [
        ("EEA 2025 provisional", eea_profile, eea_meta["rows"]),
        ("JRC technical vehicle dataset", jrc_profile, len(jrc)),
    ]:
        for concept in concepts:
            aliases = {
                "identity": ["ID", "VFN", "Mp", "Mh", "Man", "MMS", "Tan", "Va", "Ve", "Mk"] if source.startswith("EEA") else ["OEM anon", "Model anon"],
                "mass": ["mass", "m (kg)", "mt"],
                "power": ["power", "ep (kw)"],
                "fuel consumption": ["fuel consumption"],
                "electric consumption": ["electric consumption", "z (wh/km)"],
                "CO2": ["co2", "enedc", "ewltp", "ernedc", "erwltp"],
                "range": ["range"],
                "tire": ["tyre", "tire"],
                "roadload": ["f0", "f1", "f2", "roadload", "rlfi"],
                "phase": ["phase"],
            }[concept]
            fields = []
            exact_aliases = set(aliases) if concept == "identity" else set()
            for alias in aliases:
                for item in profile:
                    found = item["original_field_name"] == alias if alias in exact_aliases else alias.lower() in item["original_field_name"].lower()
                    if found and item["original_field_name"] not in fields:
                        fields.append(item["original_field_name"])
            status = "PARTIAL" if concept == "identity" and source.startswith("JRC") and fields else "DIRECT" if fields else "ABSENT"
            if concept == "roadload" and source.startswith("EEA") and fields == ["RLFI"]:
                status = "UNCLEAR"
            rows.append({
                "source": source,
                "rows": int(total),
                "engineering_concept": concept,
                "status": status,
                "direct_fields": ", ".join(fields),
                "notes": "Field presence only; no EPA semantic mapping has been applied.",
            })
    return rows


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    file_rows: list[dict[str, Any]] = []
    field_rows: list[dict[str, Any]] = []
    for path in sorted(p for p in DATA.rglob("*") if p.is_file() and "processed" not in p.parts and "staging" not in p.parts):
        file_rows.append({"source_family": source_family(path), "filename": str(path.relative_to(ROOT)), "sheet": "", "size_bytes": path.stat().st_size, "rows": "", "columns": "", "model_years": "", "notes": file_notes(path)})

    # EPA Test Car: direct row-level audit source.
    testcar_path = DATA / "raw" / "epa_testcar" / "epa_testcar_2026_raw.xlsx"
    testcar = load_excel(testcar_path, "Sheet1", 1)
    field_rows += profile_dataframe(testcar, "EPA Test Car MY2026", "Sheet1")
    grain, grain_rows, pairs, coverage = epa_audit(testcar)
    file_rows.append({"source_family": "EPA Test Car", "filename": str(testcar_path.relative_to(ROOT)), "sheet": "Sheet1", "size_bytes": testcar_path.stat().st_size, "rows": len(testcar), "columns": len(testcar.columns), "model_years": ", ".join(map(str, sorted(testcar["Model Year"].dropna().unique()))), "notes": "Header row 1; direct EPA MY2026 grain audit source."})

    # Certification workbooks: the second physical row is the header.
    for filename, label in [("light-duty-vehicle-models-2014-present.xlsx", "EPA Certified Vehicle Models"), ("light-duty-vehicle-test-results-report-2014-present.xlsx", "EPA Certified Vehicle Test Results")]:
        path = DATA / "raw" / "epa_certification" / filename
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for ws in book.worksheets:
            if ws.title == "ESRI_MAPINFO_SHEET":
                file_rows.append({"source_family": label, "filename": str(path.relative_to(ROOT)), "sheet": ws.title, "size_bytes": path.stat().st_size, "rows": 0, "columns": 0, "model_years": "", "notes": "Workbook metadata sheet; not a tabular vehicle source."})
                continue
            frame = load_excel(path, ws.title, 2)
            field_rows += profile_dataframe(frame, label, ws.title)
            years = frame.get("Model Year", pd.Series(dtype=object)).dropna().unique().tolist()
            note = "Header row 2. One row associates a certified model/carline with certification-family and standard fields; this is not treated as a complete engineering configuration." if "Models" in label else "Header row 2. Rows repeat vehicle/test context across emission results; grain is not treated as a vehicle configuration."
            file_rows.append({"source_family": label, "filename": str(path.relative_to(ROOT)), "sheet": ws.title, "size_bytes": path.stat().st_size, "rows": len(frame), "columns": len(frame.columns), "model_years": ", ".join(map(str, sorted(years)[:20])), "notes": note})
        book.close()

    # FuelEconomy workbook inside ZIP.
    fe_path = DATA / "raw" / "fueleconomy" / "fueleconomy_2026_raw.zip"
    fe_frames = load_fueleconomy(fe_path)
    for sheet, frame in fe_frames.items():
        field_rows += profile_dataframe(frame, "FuelEconomy.gov MY2026", sheet)
        year_col = next((c for c in frame.columns if str(c).lower().startswith("model y")), None)
        years = frame[year_col].dropna().unique().tolist() if year_col else []
        file_rows.append({"source_family": "FuelEconomy.gov", "filename": str(fe_path.relative_to(ROOT)), "sheet": sheet, "size_bytes": fe_path.stat().st_size, "rows": len(frame), "columns": len(frame.columns), "model_years": ", ".join(map(str, sorted(years)[:20])), "notes": "Workbook inside ZIP. Some sheets contain presentation/title rows or multiple fuel-result rows; reported as source-sheet grain, not canonical configuration grain."})

    # EEA: profile streamingly; never load or write a transformed copy.
    eea_path = DATA / "raw" / "wltp_eea" / "data.csv"
    eea_profile, eea_meta = profile_csv_stream(eea_path, "EEA 2025 provisional passenger-car data")
    field_rows += eea_profile
    file_rows.append({"source_family": "EEA WLTP / CO2 monitoring", "filename": str(eea_path.relative_to(ROOT)), "sheet": "CSV", "size_bytes": eea_path.stat().st_size, "rows": eea_meta["rows"], "columns": eea_meta["columns"], "model_years": "Reporting field 'year' present; distribution not inferred in this non-aggregating audit.", "notes": "Streaming profile. Exact null/row counts; unique cardinalities are capped above 1,000,000."})

    # JRC: first column is an unlabeled row index, retained as original source shape.
    jrc_path = DATA / "raw" / "wltp_jrc" / "Data_PV_fleet_2021_EU_PYCSIS.xlsx"
    jrc = load_excel(jrc_path, "Sheet1", 1)
    jrc.columns = ["source_row_index" if pd.isna(c) else str(c) for c in jrc.columns]
    field_rows += profile_dataframe(jrc, "JRC technical vehicle dataset", "Sheet1")
    file_rows.append({"source_family": "JRC technical vehicle dataset", "filename": str(jrc_path.relative_to(ROOT)), "sheet": "Sheet1", "size_bytes": jrc_path.stat().st_size, "rows": len(jrc), "columns": len(jrc.columns), "model_years": "No explicit model-year field", "notes": "Technical dataset with anonymized OEM/model fields and explicit WLTP/real-world roadload coefficient labels."})

    # Reference/evidence are intentionally only file/sheet inventory at this stage.
    for path in sorted((DATA / "reference").rglob("*.xlsx")):
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for ws in book.worksheets:
            file_rows.append({"source_family": source_family(path), "filename": str(path.relative_to(ROOT)), "sheet": ws.title, "size_bytes": path.stat().st_size, "rows": max(0, ws.max_row - 1), "columns": ws.max_column, "model_years": "", "notes": "Formatted reference workbook; inventoried as documentation, not ingested as data or RAG."})
        book.close()

    wltp_rows = wltp_coverage(eea_profile, profile_dataframe(jrc, "JRC technical vehicle dataset", "Sheet1"), eea_meta, jrc)
    capability = [
        {"engineering_concept": "Target ABC", "EPA Test Car": "DIRECT", "EPA Certified Test Results": "DIRECT", "FuelEconomy.gov": "ABSENT", "EEA": "ABSENT", "JRC": "ABSENT"},
        {"engineering_concept": "Set ABC", "EPA Test Car": "DIRECT", "EPA Certified Test Results": "DIRECT", "FuelEconomy.gov": "ABSENT", "EEA": "ABSENT", "JRC": "ABSENT"},
        {"engineering_concept": "Test mass / ETW", "EPA Test Car": "DIRECT", "EPA Certified Test Results": "DIRECT", "FuelEconomy.gov": "ABSENT", "EEA": "PARTIAL", "JRC": "DIRECT"},
        {"engineering_concept": "Tire specification", "EPA Test Car": "ABSENT", "EPA Certified Test Results": "ABSENT", "FuelEconomy.gov": "ABSENT", "EEA": "ABSENT", "JRC": "DIRECT"},
        {"engineering_concept": "Transmission / gears", "EPA Test Car": "DIRECT", "EPA Certified Test Results": "DIRECT", "FuelEconomy.gov": "DIRECT", "EEA": "ABSENT", "JRC": "DIRECT"},
        {"engineering_concept": "Fuel / electric consumption", "EPA Test Car": "PARTIAL", "EPA Certified Test Results": "ABSENT", "FuelEconomy.gov": "DIRECT", "EEA": "DIRECT", "JRC": "DIRECT"},
        {"engineering_concept": "CO2", "EPA Test Car": "DIRECT", "EPA Certified Test Results": "PARTIAL", "FuelEconomy.gov": "PARTIAL", "EEA": "DIRECT", "JRC": "DIRECT"},
        {"engineering_concept": "WLTP roadload coefficients", "EPA Test Car": "ABSENT", "EPA Certified Test Results": "ABSENT", "FuelEconomy.gov": "ABSENT", "EEA": "UNCLEAR", "JRC": "DIRECT"},
        {"engineering_concept": "Tire pressure / RRC", "EPA Test Car": "ABSENT", "EPA Certified Test Results": "ABSENT", "FuelEconomy.gov": "ABSENT", "EEA": "ABSENT", "JRC": "ABSENT"},
    ]
    questions = [
        {"topic": "EPA Test Car grain", "question": "Does one EPA Test Car row represent a complete test result, an emissions result, or another reporting unit when the same Make/Model/Year repeats?", "evidence": "Multiple source rows are preserved in EPA_2026_GRAIN; no aggregation was performed.", "status": "UNRESOLVED"},
        {"topic": "EPA grouping", "question": "Can Target ABC configurations ever be grouped across source rows, and if so which test identifiers/context must be equal?", "evidence": "Distinct Target/Set/ETW/procedure counts are measured, but a grouping rule is not implied.", "status": "UNRESOLVED"},
        {"topic": "EPA roadload component closure", "question": "How should the data contract represent roadload when EPA Test Car exposes Target/Set/ETW/driveline but no tire, pressure, RRC, Cd, or CdA?", "evidence": "Tier-0 coverage explicitly records these fields as absent.", "status": "ARCHITECTURE DECISION"},
        {"topic": "EEA RLFI", "question": "What is the physical meaning and unit of EEA field RLFI?", "evidence": "The local CSV exposes RLFI but no accompanying data dictionary was present.", "status": "UNRESOLVED"},
        {"topic": "JRC grain", "question": "Does each anonymized JRC row represent a real vehicle configuration, a fleet archetype, or a PYCSIS simulation input/run?", "evidence": "Anonymized OEM/model and pycsis_run fields are present; local documentation does not establish row grain.", "status": "UNRESOLVED"},
    ]
    # Keep one row per physical source file/sheet. A file-level record is retained
    # only for unstructured evidence/archive files that do not have sheet records.
    filenames_with_sheets = {row["filename"] for row in file_rows if row["sheet"]}
    file_rows = [row for row in file_rows if row["sheet"] or row["filename"] not in filenames_with_sheets]
    payload = {"generated_at": datetime.now().isoformat(timespec="seconds"), "source_files": file_rows, "field_inventory": field_rows, "epa_grain": grain, "epa_grain_rows": grain_rows, "epa_configuration_pairs": pairs, "component_coverage": coverage, "wltp_coverage": wltp_rows, "source_capability_matrix": capability, "open_questions": questions}
    (OUT / "audit_results.json").write_text(json.dumps(payload, indent=2, default=str), encoding="utf-8")
    for name, rows in [("source_files", file_rows), ("field_inventory", field_rows), ("epa_2026_grain", grain_rows), ("epa_configuration_pairs", pairs), ("component_coverage", coverage), ("wltp_coverage", wltp_rows), ("source_capability_matrix", capability), ("open_questions", questions)]:
        pd.DataFrame(rows).to_csv(OUT / f"{name}.csv", index=False, encoding="utf-8")
    print(json.dumps({"output": str(OUT.relative_to(ROOT)), "epa_grain": grain, "eea": eea_meta, "files_profiled": len(file_rows), "field_rows": len(field_rows)}, indent=2))


if __name__ == "__main__":
    main()
