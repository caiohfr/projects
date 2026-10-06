"""Render Sprint 12A report and staging workbook from the reproducible audit JSON.

Run after sprint_12a_audit.py:
    python etl/scripts/sprint_12a_deliverables.py
"""
from __future__ import annotations

import json
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[2]
AUDIT_PATH = ROOT / "etl" / "data" / "processed" / "sprint_12a_audit" / "audit_results.json"
WORKBOOK_PATH = ROOT / "etl" / "data" / "staging" / "sprint_12a_source_audit.xlsx"
REPORT_PATH = ROOT / "etl" / "reports" / "sprint_12a_source_inventory.md"
NAVY = "17365D"
BLUE = "2F75B5"
PALE = "D9EAF7"
WHITE = "FFFFFF"
THIN = Side(style="thin", color="D9E2F3")


def safe_cell(value):
    if value is None:
        return ""
    if isinstance(value, (list, tuple, dict)):
        return json.dumps(value, ensure_ascii=False)
    return value


def write_sheet(wb: Workbook, title: str, rows: list[dict], title_text: str) -> None:
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    ws["A1"] = title_text
    ws["A1"].font = Font(bold=True, color=WHITE, size=14)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(rows[0]) if rows else 1))
    if not rows:
        ws["A3"] = "No rows produced."
        return
    headers = list(rows[0].keys())
    for column, header in enumerate(headers, 1):
        cell = ws.cell(3, column, header)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN)
    for r, data in enumerate(rows, 4):
        for c, header in enumerate(headers, 1):
            cell = ws.cell(r, c, safe_cell(data.get(header, "")))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=THIN)
    last_row, last_col = 3 + len(rows), len(headers)
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(last_col)}{last_row}"
    table = Table(displayName="T" + "".join(ch for ch in title.title() if ch.isalnum()), ref=f"A3:{get_column_letter(last_col)}{last_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    ws.add_table(table)
    for c, header in enumerate(headers, 1):
        values = [str(header)] + [str(safe_cell(row.get(header, ""))) for row in rows[:500]]
        width = max(len(v) for v in values) + 2
        ws.column_dimensions[get_column_letter(c)].width = min(max(width, 12), 52)
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[3].height = 35


def build_workbook(audit: dict) -> None:
    WORKBOOK_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, "SOURCE_FILES", audit["source_files"], "Sprint 12A — Source files and sheets")
    write_sheet(wb, "FIELD_INVENTORY", audit["field_inventory"], "Sprint 12A — Field inventory (source labels retained)")
    write_sheet(wb, "SOURCE_CAPABILITY_MATRIX", audit["source_capability_matrix"], "Sprint 12A — Source capability matrix")
    write_sheet(wb, "EPA_2026_GRAIN", audit["epa_grain_rows"], "Sprint 12A — EPA MY2026 source rows (not collapsed)")
    write_sheet(wb, "EPA_CONFIGURATION_PAIRS", audit["epa_configuration_pairs"], "Sprint 12A — Candidate configuration pairs (discovery only)")
    write_sheet(wb, "COMPONENT_COVERAGE", audit["component_coverage"], "Sprint 12A — Tier-0 component coverage")
    write_sheet(wb, "WLTP_COVERAGE", audit["wltp_coverage"], "Sprint 12A — WLTP/Europe coverage (EPA semantics not imposed)")
    write_sheet(wb, "OPEN_QUESTIONS", audit["open_questions"], "Sprint 12A — Questions requiring a domain/architecture decision")
    coverage = wb["COMPONENT_COVERAGE"]
    coverage.conditional_formatting.add("G4:G100", ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B"))
    wb.save(WORKBOOK_PATH)


def pct(audit: dict, concept: str) -> float:
    return next(x["source_row_coverage_pct"] for x in audit["component_coverage"] if x["concept"] == concept)


def build_report(audit: dict) -> None:
    grain = audit["epa_grain"]
    roadload = Counter(x["roadload_availability"] for x in audit["epa_grain_rows"])
    pair_counts = Counter(x["candidate_type"] for x in audit["epa_configuration_pairs"])
    wltp = audit["wltp_coverage"]
    eea_rows = next(x["rows"] for x in wltp if x["source"].startswith("EEA"))
    jrc_rows = next(x["rows"] for x in wltp if x["source"].startswith("JRC"))
    lines = [
        "# Sprint 12A — ETL Source Inventory + Grain Audit",
        "",
        "## Scope and reproducibility",
        "",
        "This is a non-destructive source audit. Raw files were read only; no SQLite database, production ETL, application code, physics logic, RAG, enrichment, imputation, or component decomposition was created. The EPA output retains every MY2026 source row.",
        "",
        "Reproduce the analysis from the repository root:",
        "",
        "```powershell",
        "python etl/scripts/sprint_12a_audit.py",
        "python etl/scripts/sprint_12a_deliverables.py",
        "```",
        "",
        "The first command writes machine-readable audit data under `etl/data/processed/sprint_12a_audit/`; the second renders this report and `etl/data/staging/sprint_12a_source_audit.xlsx`.",
        "",
        "## Source inventory",
        "",
        f"The inventory contains {len(audit['source_files'])} source file/sheet records and {len(audit['field_inventory'])} profiled fields. It includes EPA Test Car, certification models/results, the FuelEconomy.gov MY2026 workbook inside the supplied ZIP, EEA, JRC, EV-CIS/R154 reference material, and four certification-evidence PDFs. PDFs are inventoried only.",
        "",
        "| Source | Direct observation | Grain assessment |",
        "|---|---|---|",
        "| EPA Test Car MY2026 | 3,901 source rows, 67 fields; `Sheet1` | Row grain remains unresolved; rows are preserved exactly. |",
        "| EPA Certified Vehicle Test Results | 388,420 rows, 55 fields; `Test Info` | Repeats vehicle/test context across emission-name results; not a vehicle configuration row. |",
        "| EPA Certified Vehicle Models | 26,866 rows, 10 fields; `Model Info` | Certification model/carline association; not a complete engineering configuration. |",
        "| FuelEconomy.gov MY2026 | Multiple report sheets in source ZIP | Sheet/report grain varies; presentation/title rows remain source evidence, not canonical records. |",
        f"| EEA provisional passenger-car CSV | {eea_rows:,} rows, 37 fields | Row grain cannot be confirmed from the supplied local files; no EPA mapping was imposed. |",
        f"| JRC technical dataset | {jrc_rows:,} rows, 44 fields | Anonymized OEM/model and `pycsis_run` are present; real-vehicle vs archetype/simulation grain is unresolved. |",
        "",
        "For the EEA CSV, row and null counts are exact and were measured in a single streaming pass. Field cardinalities are exact through 1,000,000 distinct values; higher cardinalities are explicitly recorded as capped lower bounds in `FIELD_INVENTORY` rather than guessed.",
        "",
        "## EPA MY2026 grain audit",
        "",
        f"- Source rows: **{grain['source_rows']:,}** (`epa_testcar_2026_raw.xlsx`, `Sheet1`).",
        f"- Make + Model + Year groups: **{grain['make_model_year_groups']:,}**; groups with multiple source rows: **{grain['groups_with_multiple_source_rows']:,}**.",
        f"- Unique observed `Test Number | Actual Tested Testgroup` combinations: **{grain['unique_test_identifiers_or_groups']:,}**.",
        f"- Distinct Target ABC sets: **{grain['distinct_target_abc_sets']:,}**; Set ABC sets: **{grain['distinct_set_abc_sets']:,}**; ETW values: **{grain['distinct_etw_variants']:,}**; procedure variants: **{grain['distinct_cycle_test_procedure_variants']:,}**.",
        f"- Groups with more than one observed Target ABC configuration: **{grain['groups_with_multiple_roadload_configurations']:,}**.",
        "",
        "Evidence: `EPA_2026_GRAIN` keeps source Excel row number, vehicle ID, configuration number, test group/number/procedure, ETW, driveline, transmission, axle/N/V, and all Target/Set A/B/C fields. This is direct observation; it is not a decision to create DB rows at that grain.",
        "",
        "### Tier-0 component coverage — EPA Test Car MY2026",
        "",
        "| Concept | Row coverage | Evidence |",
        "|---|---:|---|",
        f"| Authoritative Target ABC | {pct(audit, 'authoritative Target ABC'):.1f}% | Explicit `Target Coef A/B/C` fields |",
        f"| Set ABC | {pct(audit, 'Set ABC'):.1f}% | Explicit `Set Coef A/B/C` fields |",
        f"| ETW | {pct(audit, 'mass / ETW'):.1f}% | Explicit `Equivalent Test Weight (lbs.)` |",
        f"| Transmission / gear count / axle ratio / N/V | {pct(audit, 'transmission'):.1f}% / {pct(audit, 'gear count'):.1f}% / {pct(audit, 'axle/final-drive ratio'):.1f}% / {pct(audit, 'N/V'):.1f}% | Explicit source fields |",
        f"| Tire specification / pressure / RRC / Cd/CdA / TOTAL-NET | {pct(audit, 'tire specification'):.1f}% / {pct(audit, 'tire pressure'):.1f}% / {pct(audit, 'RRC / RR information'):.1f}% / {pct(audit, 'Cd / CdA'):.1f}% / {pct(audit, 'TOTAL / NET information'):.1f}% | No explicit Test Car field located |",
        "",
        "Roadload classification is availability-only: " + ", ".join(f"**{label}: {count:,}**" for label, count in sorted(roadload.items())) + ". No arbitrary quality threshold or component closure was applied.",
        "",
        "### Configuration-pair candidates",
        "",
        f"The conservative candidate generator wrote **{len(audit['epa_configuration_pairs']):,}** pairs: " + "; ".join(f"{name}: {count:,}" for name, count in pair_counts.items()) + ". Tire-only, mass-only, and driveline-only candidates were **0** under the strict matching rules because the Test Car table exposes no tire field and no pair met the other isolation conditions.",
        "",
        "Every pair in `EPA_CONFIGURATION_PAIRS` records both source rows, shared and differing fields, Target ABC deltas, ETW delta when numeric, warnings, and the explicit statement that it is candidate discovery—not a component causal claim.",
        "",
        "## WLTP / Europe audit",
        "",
        f"- EEA: **{eea_rows:,}** rows. Directly labelled fields cover identity/reporting keys, `m (kg)`/`Mt`, `ep (KW)`, fuel/electric consumption, NEDC/WLTP CO2 result fields, and electric range. No tire or phase field is present. `RLFI` exists but its physical meaning/unit is **UNRESOLVED** without a supplied data dictionary.",
        f"- JRC: **{jrc_rows:,}** rows. Direct fields include curb/WLTP/real-world masses, engine/electric-motor/battery attributes, tire code, gear box/gears, and explicit `wltp|f0/f1/f2` and `rw|f0/f1/f2` labels. The row's real-vehicle vs archetype/simulation meaning remains **UNRESOLVED**; anonymized identity is only partial.",
        "",
        "These findings are field-presence evidence only. No EPA terms were forced onto EEA/JRC, and no claim was made that `RLFI` or the JRC roadload fields have the same semantics as EPA Target/Set coefficients.",
        "",
        "## Reference material and evidence standard",
        "",
        "EV-CIS Data Requirements, EV-CIS Business Rules, the EV-CIS XML schema ZIP, and UNECE R154 are listed in `SOURCE_FILES` as reference documentation. They were not parsed into a RAG/vector database. The field interpretations in this audit are either **DIRECTLY OBSERVED** from explicit source labels/units and values or **UNRESOLVED**. No unsupported reference-based physical inference was elevated to confirmed.",
        "",
        "## Architectural questions for Sprint ownership",
        "",
    ]
    for i, question in enumerate(audit["open_questions"], 1):
        lines.append(f"{i}. **{question['topic']}** — {question['question']} Evidence: {question['evidence']} ({question['status']}).")
    lines += [
        "",
        "## Deliverables",
        "",
        "- `etl/notebooks/01_source_inventory.ipynb` — reproducible, non-destructive notebook entry point.",
        "- `etl/scripts/sprint_12a_audit.py` — audit implementation.",
        "- `etl/data/processed/sprint_12a_audit/` — CSV/JSON audit evidence.",
        "- `etl/data/staging/sprint_12a_source_audit.xlsx` — staging workbook with requested sheets.",
        "- `etl/reports/sprint_12a_source_inventory.md` — this closure report.",
    ]
    REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
    REPORT_PATH.write_text("\n".join(lines) + "\n", encoding="utf-8")


def verify() -> None:
    book = load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
    expected = ["SOURCE_FILES", "FIELD_INVENTORY", "SOURCE_CAPABILITY_MATRIX", "EPA_2026_GRAIN", "EPA_CONFIGURATION_PAIRS", "COMPONENT_COVERAGE", "WLTP_COVERAGE", "OPEN_QUESTIONS"]
    assert book.sheetnames == expected, book.sheetnames
    assert book["EPA_2026_GRAIN"].max_row > 3900
    assert book["FIELD_INVENTORY"].max_row > 1000
    assert book["EPA_CONFIGURATION_PAIRS"].max_row > 600
    book.close()


def main() -> None:
    audit = json.loads(AUDIT_PATH.read_text(encoding="utf-8"))
    build_workbook(audit)
    build_report(audit)
    verify()
    print(f"Wrote {WORKBOOK_PATH.relative_to(ROOT)} and {REPORT_PATH.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
